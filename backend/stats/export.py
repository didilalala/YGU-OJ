"""班级成绩单导出（Excel）：总表 + 每次作业 / 考试一个工作表。"""

import re

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from classroom.models import Assignment, ScoreRecord

from .service import class_members, class_overview, class_works, work_filter

_HEADER_FILL = PatternFill("solid", fgColor="1E2A66")
_HEADER_FONT = Font(bold=True, color="FFFFFF")
_INVALID_SHEET_CHARS = re.compile(r"[\[\]:*?/\\]")


def _sheet_title(title, used):
    """Excel 工作表名最长 31 字符，不能含 []:*?/\\，且不能重名。"""
    base = _INVALID_SHEET_CHARS.sub("_", title).strip() or "工作表"
    name, n = base[:31], 2
    while name in used:
        suffix = f"({n})"
        name, n = base[: 31 - len(suffix)] + suffix, n + 1
    used.add(name)
    return name


def _write(sheet, header, rows, widths):
    sheet.append(header)
    for cell in sheet[1]:
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    sheet.row_dimensions[1].height = 32
    for row in rows:
        sheet.append(row)
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.freeze_panes = "C2"


def build_class_workbook(class_group):
    overview = class_overview(class_group)
    works = class_works(class_group)
    workbook = openpyxl.Workbook()
    used = set()

    # 成绩总表
    summary = workbook.active
    summary.title = _sheet_title("成绩总表", used)
    header = ["学号", "姓名"]
    for row in overview["works"]:
        kind = "作业" if row["kind"] == "assignment" else "考试"
        header.append(f"{kind}：{row['title']}（{row['total_score']}）")
    header += ["总得分", "总满分", "得分率(%)", "班级排名"]
    rows = []
    for student in sorted(overview["students"], key=lambda s: s["student_no"]):
        line = [student["student_no"], student["real_name"]]
        for row in overview["works"]:
            score = student["scores"].get(f"{row['kind']}:{row['id']}")
            line.append(score if score is not None else "未提交")
        line += [student["total_score"], student["full_score"], student["rate"], student["rank"]]
        rows.append(line)
    _write(summary, header, rows, [14, 10] + [16] * len(overview["works"]) + [10, 10, 10, 10])

    # 每次作业 / 考试的各题得分
    members = class_members(class_group)
    member_ids = [user.id for user, _, _ in members]
    for work in works:
        is_assignment = isinstance(work, Assignment)
        sheet = workbook.create_sheet(_sheet_title(("作业-" if is_assignment else "考试-") + work.title, used))
        items = list(work.items.select_related("problem"))
        records = {r.student_id: r for r in ScoreRecord.objects.filter(student_id__in=member_ids, **work_filter(work))}
        header = ["学号", "姓名"] + [f"{i + 1}. {item.problem.title}（{item.score}）" for i, item in enumerate(items)]
        header += ["总分", "通过题数"] + (["含迟交"] if is_assignment else [])
        rows = []
        for user, no, name in members:
            record = records.get(user.id)
            if record is None:
                rows.append([no, name] + ["" for _ in items] + ["未提交", 0] + ([""] if is_assignment else []))
                continue
            line = [no, name]
            for item in items:
                entry = record.details.get(str(item.problem_id))
                line.append(entry["score"] if entry else "")
            line += [record.score, record.solved_count] + (["是" if record.is_late else ""] if is_assignment else [])
            rows.append(line)
        _write(sheet, header, rows, [14, 10] + [14] * len(items) + [10, 10, 8])

    return workbook
