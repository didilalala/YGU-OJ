"""
学生名单导入（设计说明书 9.1）。

1. 逐行校验学号格式、文件内重复；
2. 学号已有账号的直接加入班级，新学号创建账号（初始密码为学号后六位，首次登录强制改密）；
3. 返回导入报告：新建账号数、加入已有账号数、失败行号及原因。

校验通过的行在一个事务中导入，失败的行不影响其他行。
"""

import zipfile
from dataclasses import dataclass, field
from datetime import date, datetime

import openpyxl
from django.db import transaction
from django.utils import timezone
from openpyxl.styles import Font, PatternFill
from openpyxl.utils.exceptions import InvalidFileException

from account.models import StudentProfile, User
from account.students import is_valid_student_no, set_initial_password

from stats.cache import bump_class

from .models import ClassMember

# 表头 → 字段；表头中的空格与必填标记 * 会被忽略
COLUMNS = {"学号": "student_no", "姓名": "real_name", "学院": "college", "专业": "major", "年级": "grade"}
LABELS = {field: title for title, field in COLUMNS.items()}
REQUIRED_COLUMNS = ("学号", "姓名")
MAX_LENGTHS = {"real_name": 32, "college": 64, "major": 64, "grade": 16}
MAX_ROWS = 1000
MAX_FILE_SIZE = 5 * 1024 * 1024


class StudentImportError(ValueError):
    """名单文件整体无法读取，消息可直接展示给教师。"""


@dataclass
class StudentRow:
    row: int | None  # Excel 行号；单个添加时为 None
    student_no: str = ""
    real_name: str = ""
    college: str = ""
    major: str = ""
    grade: str = ""


@dataclass
class ImportReport:
    created: int = 0  # 新建账号并加入
    joined: int = 0  # 已有账号加入（或重新加入）
    already_in_class: int = 0  # 已在班级中，无需处理
    failed: list = field(default_factory=list)
    notes: list = field(default_factory=list)

    def fail(self, row, reason):
        self.failed.append({"row": row.row, "student_no": row.student_no, "reason": reason})

    def as_dict(self):
        return {
            "created": self.created,
            "joined": self.joined,
            "already_in_class": self.already_in_class,
            "failed": self.failed,
            "notes": self.notes,
        }


def _cell(value):
    if value is None:
        return ""
    # 学号列常被 Excel 存成数字，如 20260001 读出为 20260001 或 20260001.0
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, (datetime, date)):
        return value.strftime("%Y")
    return str(value).strip()


def read_rows(file):
    """读取第一个工作表；按表头名称识别列，列顺序不限。"""
    try:
        workbook = openpyxl.load_workbook(file, read_only=True, data_only=True)
    except (InvalidFileException, zipfile.BadZipFile, KeyError, OSError, ValueError, TypeError) as exc:
        raise StudentImportError("无法读取文件，请上传 .xlsx 格式的名单（可下载模板填写）") from exc

    try:
        rows = workbook.worksheets[0].iter_rows(values_only=True)
        header = next(rows, None) or ()
        positions = {}
        for index, title in enumerate(header):
            name = _cell(title).replace("*", "").replace(" ", "")
            if name in COLUMNS:
                positions[COLUMNS[name]] = index
        missing = [title for title in REQUIRED_COLUMNS if COLUMNS[title] not in positions]
        if missing:
            raise StudentImportError(f"第一行表头缺少列：{'、'.join(missing)}")

        result = []
        for line, values in enumerate(rows, start=2):
            data = {key: _cell(values[i]) if i < len(values) else "" for key, i in positions.items()}
            if not any(data.values()):
                continue  # 跳过空行
            result.append(StudentRow(row=line, **data))
            if len(result) > MAX_ROWS:
                raise StudentImportError(f"一次最多导入 {MAX_ROWS} 名学生")
    finally:
        workbook.close()

    if not result:
        raise StudentImportError("名单中没有学生")
    return result


def validate_rows(rows, report):
    """校验各行，返回通过校验的行；失败的行记入报告。"""
    valid, seen = [], {}
    for row in rows:
        if not row.student_no:
            report.fail(row, "学号为空")
        elif not is_valid_student_no(row.student_no):
            report.fail(row, "学号格式不正确")
        elif not row.real_name:
            report.fail(row, "姓名为空")
        elif row.student_no in seen:
            report.fail(row, f"与第 {seen[row.student_no]} 行学号重复")
        elif too_long := next((k for k, limit in MAX_LENGTHS.items() if len(getattr(row, k)) > limit), None):
            report.fail(row, f"{LABELS[too_long]}过长（最多 {MAX_LENGTHS[too_long]} 字）")
        else:
            seen[row.student_no] = row.row
            valid.append(row)
    return valid


def _join(class_group, user, report):
    """加入班级；返回是否为本次新加入（含重新加入已移出的学生）。"""
    member = ClassMember.objects.filter(class_group=class_group, student=user).first()
    if member is None:
        ClassMember.objects.create(class_group=class_group, student=user)
        return True
    if member.is_removed:
        member.is_removed = False
        member.removed_at = None
        member.save(update_fields=["is_removed", "removed_at"])
        return True
    report.already_in_class += 1
    return False


@transaction.atomic
def import_students(class_group, rows):
    report = ImportReport()
    bump_class(class_group.id)
    valid = validate_rows(rows, report)
    student_nos = [row.student_no for row in valid]
    profiles = {
        p.student_no: p for p in StudentProfile.objects.filter(student_no__in=student_nos).select_related("user")
    }
    taken_usernames = set(
        User.objects.filter(username__in=set(student_nos) - set(profiles)).values_list("username", flat=True)
    )

    for row in valid:
        profile = profiles.get(row.student_no)
        if profile is not None:
            user = profile.user
            if not user.is_active:
                report.fail(row, "该学号的账号已停用")
                continue
            if _join(class_group, user, report):
                report.joined += 1
            if profile.real_name != row.real_name:
                report.notes.append(
                    {
                        "row": row.row,
                        "student_no": row.student_no,
                        "message": f"姓名与已有账号不一致（已有：{profile.real_name}），未修改已有信息",
                    }
                )
            continue

        if row.student_no in taken_usernames:
            # 如普通用户注册时恰好用了这个学号作用户名
            report.fail(row, "该学号已被其他账号用作用户名，请联系管理员处理")
            continue

        user = User(username=row.student_no, role=User.Role.STUDENT, source=User.Source.IMPORT)
        set_initial_password(user, row.student_no)
        user.save()
        StudentProfile.objects.create(
            user=user,
            student_no=row.student_no,
            real_name=row.real_name,
            college=row.college,
            major=row.major,
            grade=row.grade,
        )
        ClassMember.objects.create(class_group=class_group, student=user)
        report.created += 1

    return report


def remove_members(class_group, user_ids):
    """移出班级（保留历史），返回实际移出的人数。"""
    bump_class(class_group.id)
    return ClassMember.objects.active().filter(class_group=class_group, student_id__in=user_ids).update(
        is_removed=True, removed_at=timezone.now()
    )


def reset_passwords(class_group, user_ids):
    """把本班学生的密码重置为初始密码，再次登录后强制修改；返回重置人数。"""
    members = (
        ClassMember.objects.active()
        .filter(class_group=class_group, student_id__in=user_ids)
        .select_related("student__student_profile")
    )
    count = 0
    for member in members:
        user = member.student
        profile = getattr(user, "student_profile", None)
        if profile is None:
            continue
        set_initial_password(user, profile.student_no)
        user.save(update_fields=["password", "is_initial_password"])
        count += 1
    return count


def build_template():
    """导入模板：第一个工作表只有表头，填写说明放在第二个工作表，避免示例行被误导入。"""
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "学生名单"
    headers = ["学号*", "姓名*", "学院", "专业", "年级"]
    sheet.append(headers)
    fill = PatternFill("solid", fgColor="1E2A66")
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = fill
    for column, width in zip("ABCDE", (18, 12, 20, 20, 10)):
        sheet.column_dimensions[column].width = width
    # 学号列设为文本格式，避免长学号被 Excel 显示为科学计数法（设在列上，不产生空行）
    sheet.column_dimensions["A"].number_format = "@"

    notes = workbook.create_sheet("填写说明")
    for line in (
        "1. 在「学生名单」工作表中从第 2 行开始填写，每行一名学生；学号、姓名必填。",
        "2. 学号已有账号的学生直接加入班级，不修改其已有信息。",
        "3. 新学号自动创建账号：用户名为学号，初始密码为学号后六位，首次登录须修改密码。",
        f"4. 一次最多导入 {MAX_ROWS} 名学生。",
    ):
        notes.append([line])
    notes.column_dimensions["A"].width = 80
    return workbook
