import io
import time

import openpyxl
from django.contrib.auth import authenticate
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import override_settings
from rest_framework.test import APIClient, APITestCase

from account.models import StudentProfile, User

from .models import ClassGroup, ClassMember

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def make_xlsx(rows, header=("学号*", "姓名*", "学院", "专业", "年级")):
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    if header:
        sheet.append(list(header))
    for row in rows:
        sheet.append(list(row))
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


class ClassTestBase(APITestCase):
    def setUp(self):
        self.teacher = User.objects.create_user("t1", password="x", role=User.Role.TEACHER)
        self.other_teacher = User.objects.create_user("t2", password="x", role=User.Role.TEACHER)
        self.admin = User.objects.create_user("admin1", password="x", role=User.Role.ADMIN)
        self.klass = ClassGroup.objects.create(name="软件 2601", course="C 语言", semester="2026-2027-1", teacher=self.teacher)

    def client_for(self, user):
        client = APIClient()
        client.force_authenticate(user)
        return client

    def make_student(self, student_no, name="已有学生", **fields):
        user = User.objects.create_user(student_no, password="Pass-2026!", role=User.Role.STUDENT, **fields)
        StudentProfile.objects.create(user=user, student_no=student_no, real_name=name)
        return user

    def upload(self, content, user=None, class_id=None, name="名单.xlsx"):
        file = SimpleUploadedFile(name, content, XLSX)
        return self.client_for(user or self.teacher).post(
            f"/api/admin/classes/{class_id or self.klass.id}/import", {"file": file}, format="multipart"
        )


class ImportTests(ClassTestBase):
    def test_import_creates_accounts_and_joins_existing(self):
        existing = self.make_student("20260099", name="王五")
        content = make_xlsx(
            [
                (20260001, "张三", "计算机学院", "软件工程", 2026),  # 学号被 Excel 存为数字
                ("20260002", "李四", None, None, None),
                ("20260099", "王五（改名）", None, None, None),
            ]
        )
        response = self.upload(content)

        self.assertEqual(response.status_code, 200)
        report = response.json()["data"]
        self.assertEqual((report["created"], report["joined"], report["already_in_class"]), (2, 1, 0))
        self.assertEqual(report["failed"], [])
        self.assertEqual(report["notes"][0]["row"], 4)
        self.assertIn("姓名与已有账号不一致", report["notes"][0]["message"])

        user = User.objects.get(username="20260001")
        self.assertEqual((user.role, user.source, user.is_initial_password), ("student", "import", True))
        self.assertEqual(user.student_profile.real_name, "张三")
        self.assertEqual(user.student_profile.grade, "2026")
        # 已有学生的信息不被覆盖
        existing.student_profile.refresh_from_db()
        self.assertEqual(existing.student_profile.real_name, "王五")
        self.assertEqual(ClassMember.objects.filter(class_group=self.klass).count(), 3)

    def test_initial_password_is_last_six_digits_and_upgraded_on_login(self):
        self.upload(make_xlsx([("20260001", "张三"), ("1234", "短学号")]))

        user = authenticate(username="20260001", password="260001")
        self.assertIsNotNone(user)
        self.assertTrue(User.objects.get(username="1234").check_password("1234"))
        # 首次登录后自动改用默认的强哈希
        self.assertTrue(User.objects.get(username="20260001").password.startswith("pbkdf2_sha256$"))
        self.assertTrue(User.objects.get(username="1234").password.startswith("pbkdf2_sha256$"))

    def test_row_level_failures_do_not_block_other_rows(self):
        User.objects.create_user("20260005", password="x")  # 普通用户占用了这个用户名
        self.make_student("20260006").__class__.objects.filter(username="20260006").update(is_active=False)
        content = make_xlsx(
            [
                ("20260001", "张三"),
                ("", "无学号"),
                ("abc", "格式错"),
                ("20260003", ""),
                ("20260001", "重复"),
                ("20260004", "名" * 33),
                ("20260005", "被占用"),
                ("20260006", "已停用"),
                (None, None, None, None, None),  # 空行跳过
            ]
        )
        report = self.upload(content).json()["data"]

        self.assertEqual(report["created"], 1)
        self.assertEqual(
            [(f["row"], f["reason"]) for f in report["failed"]],
            [
                (3, "学号为空"),
                (4, "学号格式不正确"),
                (5, "姓名为空"),
                (6, "与第 2 行学号重复"),
                (7, "姓名过长（最多 32 字）"),
                (8, "该学号已被其他账号用作用户名，请联系管理员处理"),
                (9, "该学号的账号已停用"),
            ],
        )

    def test_reimport_and_rejoin_removed_students(self):
        self.upload(make_xlsx([("20260001", "张三"), ("20260002", "李四")]))
        student = User.objects.get(username="20260001")
        self.client_for(self.teacher).post(
            f"/api/admin/classes/{self.klass.id}/remove-students", {"user_ids": [student.id]}, format="json"
        )
        report = self.upload(make_xlsx([("20260001", "张三"), ("20260002", "李四")])).json()["data"]

        self.assertEqual((report["created"], report["joined"], report["already_in_class"]), (0, 1, 1))
        self.assertFalse(ClassMember.objects.get(student=student).is_removed)
        self.assertEqual(ClassMember.objects.filter(class_group=self.klass).count(), 2)

    def test_same_student_in_multiple_classes(self):
        other = ClassGroup.objects.create(name="软件 2602", teacher=self.teacher)
        self.upload(make_xlsx([("20260001", "张三")]))
        report = self.upload(make_xlsx([("20260001", "张三")]), class_id=other.id).json()["data"]

        self.assertEqual(report["joined"], 1)
        self.assertEqual(User.objects.filter(username="20260001").count(), 1)

    def test_columns_matched_by_header_name(self):
        content = make_xlsx([("计算机学院", "张三", "20260001")], header=("学院", "姓名", "学号"))
        self.assertEqual(self.upload(content).json()["data"]["created"], 1)
        self.assertEqual(StudentProfile.objects.get(student_no="20260001").college, "计算机学院")

    def test_file_level_errors(self):
        cases = [
            (b"not an excel file", "无法读取文件"),
            (make_xlsx([("20260001",)], header=("学号",)), "表头缺少列：姓名"),
            (make_xlsx([], header=("学号", "姓名")), "名单中没有学生"),
        ]
        for content, message in cases:
            with self.subTest(message=message):
                response = self.upload(content)
                self.assertEqual(response.status_code, 400)
                self.assertIn(message, response.json()["message"])

    @override_settings(STUDENT_NO_PATTERN=r"^[A-Z]\d{6}$")
    def test_student_no_pattern_is_configurable(self):
        report = self.upload(make_xlsx([("S260001", "张三"), ("20260001", "李四")])).json()["data"]
        self.assertEqual(report["created"], 1)
        self.assertEqual(report["failed"][0]["reason"], "学号格式不正确")

    def test_import_is_fast_enough_for_a_large_class(self):
        rows = [(f"2026{i:04d}", f"学生{i}") for i in range(300)]
        start = time.time()
        report = self.upload(make_xlsx(rows)).json()["data"]
        self.assertEqual(report["created"], 300)
        self.assertLess(time.time() - start, 30)

    def test_other_teachers_cannot_import(self):
        response = self.upload(make_xlsx([("20260001", "张三")]), user=self.other_teacher)
        self.assertEqual(response.status_code, 404)
        self.assertEqual(self.upload(make_xlsx([("20260001", "张三")]), user=self.admin).status_code, 200)

    def test_template_download(self):
        response = self.client_for(self.teacher).get("/api/admin/classes/import-template")

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response["Content-Type"], XLSX)
        self.assertIn("attachment", response["Content-Disposition"])
        workbook = openpyxl.load_workbook(io.BytesIO(response.content))
        self.assertEqual([c.value for c in workbook.worksheets[0][1]], ["学号*", "姓名*", "学院", "专业", "年级"])
        self.assertEqual(workbook.worksheets[0].max_row, 1)  # 没有示例行
        # 模板本身可以直接导入（只有表头时提示没有学生）
        self.assertIn("名单中没有学生", self.upload(response.content).json()["message"])


class ClassApiTests(ClassTestBase):
    def test_teacher_sees_only_own_classes(self):
        ClassGroup.objects.create(name="别人的班", teacher=self.other_teacher)
        data = self.client_for(self.teacher).get("/api/admin/classes").json()["data"]

        self.assertEqual([c["name"] for c in data["results"]], ["软件 2601"])
        self.assertEqual(self.client_for(self.admin).get("/api/admin/classes").json()["data"]["total"], 2)

    def test_create_update_delete(self):
        client = self.client_for(self.teacher)
        response = client.post("/api/admin/classes", {"name": "新班", "course": "Java", "semester": "2026-2027-2"}, format="json")
        self.assertEqual(response.status_code, 201)
        data = response.json()["data"]
        self.assertEqual((data["teacher"], data["member_count"], len(data["invite_code"])), (self.teacher.id, 0, 8))

        self.assertEqual(client.patch(f"/api/admin/classes/{data['id']}", {"name": "改名"}, format="json").status_code, 200)
        self.assertEqual(client.delete(f"/api/admin/classes/{data['id']}").status_code, 204)

    def test_delete_keeps_student_accounts(self):
        self.upload(make_xlsx([("20260001", "张三")]))
        self.client_for(self.teacher).delete(f"/api/admin/classes/{self.klass.id}")
        self.assertTrue(User.objects.filter(username="20260001").exists())

    def test_other_teacher_cannot_access_class(self):
        client = self.client_for(self.other_teacher)
        self.assertEqual(client.get(f"/api/admin/classes/{self.klass.id}").status_code, 404)
        self.assertEqual(client.get(f"/api/admin/classes/{self.klass.id}/students").status_code, 404)
        self.assertEqual(client.patch(f"/api/admin/classes/{self.klass.id}", {"name": "x"}, format="json").status_code, 404)

    def test_students_cannot_use_admin_api(self):
        student = self.make_student("20260001")
        self.assertEqual(self.client_for(student).get("/api/admin/classes").status_code, 403)

    def test_student_list_member_count_and_filters(self):
        self.upload(make_xlsx([("20260002", "李四"), ("20260001", "张三"), ("20260003", "王五")]))
        removed = User.objects.get(username="20260003")
        client = self.client_for(self.teacher)
        client.post(f"/api/admin/classes/{self.klass.id}/remove-students", {"user_ids": [removed.id]}, format="json")

        data = client.get(f"/api/admin/classes/{self.klass.id}/students").json()["data"]
        self.assertEqual([s["student_no"] for s in data["results"]], ["20260001", "20260002"])
        self.assertTrue(data["results"][0]["is_initial_password"])
        all_members = client.get(f"/api/admin/classes/{self.klass.id}/students?include_removed=1").json()["data"]
        self.assertEqual(all_members["total"], 3)
        self.assertTrue(all_members["results"][-1]["is_removed"])
        self.assertEqual(client.get(f"/api/admin/classes/{self.klass.id}/students?keyword=李").json()["data"]["total"], 1)
        self.assertEqual(client.get(f"/api/admin/classes/{self.klass.id}").json()["data"]["member_count"], 2)

    def test_add_single_student(self):
        client = self.client_for(self.teacher)
        url = f"/api/admin/classes/{self.klass.id}/students"
        response = client.post(url, {"student_no": "20260001", "real_name": "张三", "major": "软件工程"}, format="json")
        self.assertEqual(response.status_code, 201)
        self.assertEqual(response.json()["data"]["created"], 1)

        response = client.post(url, {"student_no": "20260001", "real_name": "张三"}, format="json")
        self.assertEqual(response.status_code, 400)
        self.assertEqual(response.json()["message"], "该学生已在班级中")
        response = client.post(url, {"student_no": "x1", "real_name": "张三"}, format="json")
        self.assertIn("学号格式不正确", response.json()["message"])

    def test_reset_password_only_for_own_class_members(self):
        self.upload(make_xlsx([("20260001", "张三")]))
        student = User.objects.get(username="20260001")
        student.set_password("Changed-2026!")
        student.is_initial_password = False
        student.save()
        outsider = self.make_student("20269999")

        response = self.client_for(self.teacher).post(
            f"/api/admin/classes/{self.klass.id}/reset-password",
            {"user_ids": [student.id, outsider.id]},
            format="json",
        )
        self.assertEqual(response.json()["data"], {"count": 1})
        student.refresh_from_db()
        outsider.refresh_from_db()
        self.assertTrue(student.check_password("260001"))
        self.assertTrue(student.is_initial_password)
        self.assertTrue(outsider.check_password("Pass-2026!"))

        response = self.client_for(self.other_teacher).post(
            f"/api/admin/classes/{self.klass.id}/reset-password", {"user_ids": [student.id]}, format="json"
        )
        self.assertEqual(response.status_code, 404)

    def test_user_ids_required(self):
        response = self.client_for(self.teacher).post(
            f"/api/admin/classes/{self.klass.id}/remove-students", {"user_ids": []}, format="json"
        )
        self.assertEqual(response.status_code, 400)


class MyClassTests(ClassTestBase):
    def test_student_sees_active_classes(self):
        student = self.make_student("20260001")
        other = ClassGroup.objects.create(name="已移出的班", teacher=self.other_teacher)
        ClassMember.objects.create(class_group=self.klass, student=student)
        ClassMember.objects.create(class_group=other, student=student, is_removed=True)
        StudentProfile.objects.filter(user=self.teacher).delete()

        data = self.client_for(student).get("/api/classes/mine").json()["data"]
        self.assertEqual(data, [{
            "id": self.klass.id, "name": "软件 2601", "course": "C 语言", "semester": "2026-2027-1", "teacher_name": "t1",
        }])

    def test_anonymous(self):
        self.assertEqual(APIClient().get("/api/classes/mine").status_code, 401)
